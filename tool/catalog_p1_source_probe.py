#!/usr/bin/env python3
"""Read-only official education/completions and occupational projections probe.

Runs only on GitHub Actions' ephemeral runner. Prints aggregate diagnostics and
column names, NEVER downloads files into git or prints individual learner rows.
"""
import csv
import hashlib
import io
import os
import sys
import urllib.error
import urllib.request
import zipfile
from pathlib import Path

URLS = {
    "IPEDS 2025 provisional awards": "https://nces.ed.gov/ipeds/datacenter/data/C2025_A.zip",
    "BLS 2025-2035 occupation workbook": "https://www.bls.gov/emp/ind-occ-matrix/occupation.xlsx",
}
def fetch(name, url):
    req = urllib.request.Request(url,headers={"User-Agent":"Mozilla/5.0 (compatible; LearningAppCatalogEvidence/1.0; admin@invalid.example)","Accept":"*/*"})
    with urllib.request.urlopen(req, timeout=90) as r:
        data = r.read(60*1024*1024+1)
        ctype=r.headers.get('Content-Type')
    if len(data)>60*1024*1024: raise ValueError("Input exceeds max allowed download size")
    print(name, "bytes",len(data),"SHA256",hashlib.sha256(data).hexdigest(),"Content-Type",ctype,flush=True)
    return data

for name,url in URLS.items():
    try:
        data=fetch(name,url)
        if data[:2]==b'PK':
            z=zipfile.ZipFile(io.BytesIO(data))
            names=z.namelist()
            print("ZIP members:",names[:10],flush=True)
            if "IPEDS" in name:
                item=next((n for n in names if n.upper().endswith(".CSV") and "C2025_A" in n.upper()),None)
                if item:
                    with z.open(item) as f:
                        txt=io.TextIOWrapper(f,encoding="utf-8-sig",errors="replace",newline="")
                        reader=csv.reader(txt)
                        print("IPEDS header:",next(reader),flush=True)
                        print("IPEDS sample with limited non-identifying cols",[[str(v)[:22] for v in row[:7]] for _,row in zip(range(2),reader)],flush=True)
            else:
                import openpyxl
                w=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True)
                print("Workbooks",w.sheetnames,flush=True)
                for sheet in w.worksheets[:3]:
                    print("Sheet",sheet.title,"size",sheet.max_row,sheet.max_column,flush=True)
                    for i,row in enumerate(sheet.iter_rows(min_row=1,max_row=12,values_only=True)):
                        print("   row",i,repr(tuple(str(v)[:60] for v in row[:18])),flush=True)
        else:
            print("Unexpected input signature",repr(data[:60]),flush=True)
    except Exception as e:
        print("PROBE FAILURE",name,type(e).__name__,str(e)[:230],flush=True)
        sys.exitcode=2
