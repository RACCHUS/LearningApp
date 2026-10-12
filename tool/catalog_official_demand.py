#!/usr/bin/env python3
"""Source-pinned, review-only national IPEDS/BLS demand intake.

Usage:
  python tool/catalog_official_demand.py bls --validate
  python tool/catalog_official_demand.py bls --import --confirm-project xzvkdwebtbxlrxagtzlv
  python tool/catalog_official_demand.py ipeds --file /tmp/C2025_A.zip --source-url https://... --validate
A published CIP-SOC relationship is NOT a license/certification popularity score.
"""
from __future__ import annotations
import argparse
import csv
import hashlib
import io
import json
import math
import os
import re
import sys
import urllib.error
import urllib.request
import zipfile
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

PROJECT = "xzvkdwebtbxlrxagtzlv"
BLS_URL = "https://www.bls.gov/emp/ind-occ-matrix/occupation.xlsx"
BLS_SHA256 = "a7d060f0576f7d4f5b829b9a0bb4a6c11a12766d159ebb0e07aab5bd2e7583c5"
IPEDS_INFO = "https://nces.ed.gov/ipeds/datacenter/DataFiles.aspx?fromIpeds=true&gotoReportId=7"
MAX_BYTES = 60 * 1024 * 1024
CIP = re.compile(r"^\d{2}\.\d{4}$")
SOC = re.compile(r"^\d{2}-\d{4}$")
AGGREGATE_TYPES = {"Summary", "Total"}

def fetch(url: str, cap: int = MAX_BYTES) -> bytes:
    request = urllib.request.Request(
        url, headers={"User-Agent": "Mozilla/5.0 LearningAppDemandEvidence/1.0",
                      "Accept": "application/octet-stream,application/zip,*/*"}
    )
    with urllib.request.urlopen(request, timeout=120) as response:
        chunks = []
        count = 0
        while True:
            chunk = response.read(256 * 1024)
            if not chunk: break
            count += len(chunk)
            if count > cap: raise ValueError("Source exceeds maximum allowed size")
            chunks.append(chunk)
    return b"".join(chunks)

def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()

def exact_positive_decimal(raw: object) -> Decimal | None:
    if raw is None or str(raw).strip() in ("", "—", "-", "NA", "N/A"): return None
    try:
        value = Decimal(str(raw).strip().replace(",", ""))
        if not value.is_finite() or value < 0: return None
        return value
    except InvalidOperation:
        return None

def make_row(system, version, code, metric, value, unit, period, dimension,
             publisher, release, url, digest, method):
    return {"taxonomy_system": system, "taxonomy_version": version,
            "taxonomy_code": code, "metric": metric, "metric_value": str(value),
            "metric_unit": unit, "period": period, "geography": "US",
            "dimension": dimension, "publisher": publisher,
            "source_release": release, "source_url": url, "source_sha256": digest,
            "method_note": method}

def parse_bls(blob: bytes) -> list[dict]:
    # Data contain both nested Summary rows and Line item rows. Counting both
    # would double-count US occupations. Values are *thousands* of jobs.
    import openpyxl
    if blob[:2] != b"PK": raise ValueError("BLS source is not an XLSX workbook")
    book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    if "Table 1.2" not in book.sheetnames: raise ValueError("Official BLS Table 1.2 missing")
    sheet = book["Table 1.2"]
    iterator = sheet.iter_rows(values_only=True)
    heading = next(iterator)
    if "2025" not in str(heading[0]) or "2035" not in str(heading[0]):
        raise ValueError("BLS projection window changed: update reviewed parser")
    cols = [str(v or "").strip() for v in next(iterator)]
    required = ["2025 National Employment Matrix code", "Occupation type",
                "Occupational openings, 2025–35 annual average"]
    if not all(field in cols for field in required):
        raise ValueError(f"BLS header changed; missing {set(required)-set(cols)}")
    ix_code, ix_kind, ix_openings = (cols.index(x) for x in required)
    if "thousands" not in str(heading[0]).lower():
        raise ValueError("Cannot verify BLS openings are in thousands")
    result = []
    codes = set()
    digest = sha256(blob)
    for record in iterator:
        code, kind = str(record[ix_code] or "").strip(), str(record[ix_kind] or "").strip()
        if kind != "Line item" or not SOC.fullmatch(code): continue
        openings = exact_positive_decimal(record[ix_openings])
        if openings is None: continue
        if openings > 10000: raise ValueError("Implausible BLS openings (thousands)")
        if code in codes: raise ValueError("Duplicate BLS detailed SOC code")
        codes.add(code)
        result.append(make_row(
            "bls_soc", "soc_2018", code, "bls_annual_openings_thousands",
            openings, "thousands_of_jobs", "2025-2035", "all",
            "BLS_EMP", "BLS_EMP_2025_2035", BLS_URL, digest,
            "BLS Table 1.2 2025-2035 line-item SOC; average annual occupational openings in thousands, rounded as published; no summary rows."
        ))
    if len(result) < 700:
        raise ValueError(f"BLS detailed SOC count too small: {len(result)}")
    return result

def parse_ipeds(blob: bytes, source_url: str) -> list[dict]:
    # Only an operator-verified official ZIP passed via --file is accepted.
    # There is intentionally no guessed C2025_A ZIP URL: 2026 official NCES
    # data center uses release links, and a conventional static URL returned 404.
    if blob[:2] != b"PK": raise ValueError("IPEDS source is not a ZIP")
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        entries = [n for n in z.namelist()
                   if n.upper().endswith(".CSV") and "C2025_A" in Path(n).name.upper()]
        if len(entries) != 1: raise ValueError("ZIP must contain exactly one C2025_A CSV")
        entry = z.getinfo(entries[0])
        if entry.file_size > 160 * 1024 * 1024:
            raise ValueError("IPEDS CSV too large to parse safely")
        aggregate = defaultdict(int)
        valid = 0
        with z.open(entry) as fp:
            reader = csv.DictReader(io.TextIOWrapper(fp, encoding="utf-8-sig", newline=""))
            required = {"UNITID","CIPCODE","AWLEVEL","CTOTALT"}
            cols = set(reader.fieldnames or [])
            if not required <= cols:
                raise ValueError(f"IPEDS C2025_A dictionary headers missing: {required-cols}")
            for row in reader:
                code = (row["CIPCODE"] or "").strip()
                award_level = (row["AWLEVEL"] or "").strip()
                if not CIP.fullmatch(code) or not award_level.isdigit(): continue
                count = (row["CTOTALT"] or "").strip()
                if not count.isdecimal(): continue
                n = int(count)
                # NCES special/suppression negatives and subtotal code rows
                # were excluded above. Do not aggregate '99' CIP categories.
                if n > 1_000_000: raise ValueError("Implausible awards on single row")
                aggregate[(code, award_level)] += n
                valid += 1
    if valid < 5000 or len(aggregate) < 1000:
        raise ValueError(f"Too few validated IPEDS awards rows: {valid}, groups {len(aggregate)}")
    digest = sha256(blob)
    return [make_row(
        "cip","2020",code,"ipeds_awards_conferred",n,"awards",
        "2024-2025",f"awardlevel:{level}","NCES_IPEDS",
        "C2025_A_provisional",source_url,digest,
        "IPEDS provisional C2025_A: sum institutional CTOTALT awards for an exact 6-digit CIP and AWLEVEL, July 1 2024 to June 30 2025; not unique students, enrollment, courses or provider openings."
    ) for (code,level),n in sorted(aggregate.items())]

def request_json(url: str, key: str, method: str="GET",
                 payload: list[dict] | dict | None=None, extra: dict | None=None):
    headers = {"apikey": key, "Accept":"application/json",
               "User-Agent":"LearningAppCatalogDemand/1.0"}
    if payload is not None:
        headers.update({"Content-Type":"application/json",
                        "Prefer":"resolution=ignore-duplicates,return=minimal"})
    if extra: headers.update(extra)
    data = None if payload is None else json.dumps(payload).encode()
    request = urllib.request.Request(url, method=method, data=data, headers=headers)
    try:
        with urllib.request.urlopen(request, timeout=90) as r:
            response = r.read(10_000)
            return json.loads(response) if response else None
    except urllib.error.HTTPError as err:
        # Deliberately do not echo response/body or privileged URL.
        raise RuntimeError(f"Supabase API HTTP {err.code} during {method}") from err

def ingest(rows: list[dict], source_url: str, digest: str,
           size: int, project_ref: str, release: str):
    if project_ref != PROJECT: raise ValueError("Wrong target project")
    root = os.environ.get("SUPABASE_URL", "").rstrip("/")
    if root != f"https://{PROJECT}.supabase.co":
        raise ValueError("Wrong SUPABASE_URL; refusing write")
    key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "").strip()
    if not key.startswith("sb_secret_") or len(key) < 20:
        raise ValueError("Modern sb_secret_ service API key required")
    # Do not rewrite an existing pinned release with different content.
    prior = request_json(
        f"{root}/rest/v1/catalog_official_demand_imports?select=sha256,observations&source_release=eq.{release}",
        key)
    if prior and prior[0]["sha256"] != digest:
        raise ValueError("Previously loaded release SHA changed; requires reviewed new release")
    observed = {row["taxonomy_code"] for row in rows}
    if len(observed) < 700:
        raise ValueError("Insufficient distinct taxonomy codes before live write")
    for offset in range(0,len(rows),200):
        request_json(
            f"{root}/rest/v1/catalog_official_demand_observations?on_conflict=taxonomy_system,taxonomy_version,taxonomy_code,metric,period,geography,dimension,source_release",
            key, method="POST", payload=rows[offset:offset+200])
    if not prior:
        request_json(f"{root}/rest/v1/catalog_official_demand_imports",key,
                     method="POST",payload=[{
                        "source_release":release,"source_url":source_url,
                        "sha256":digest,"source_bytes":size,"observations":len(rows),
                        "method_note":"Official sourced aggregate, private until reviewed. Checksummed immutable release; import uses no overwrite."
                     }])
    print(json.dumps({"result":"submitted_private_sourced_evidence",
                      "source_release":release,"rows_attempted":len(rows),
                      "distinct_codes":len(observed),
                      "sha256":digest,"public_credential_records_changed":0},indent=2))

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source",choices=["bls","ipeds"])
    parser.add_argument("--file",help="Operator-provided original downloaded official source (required for IPEDS)")
    parser.add_argument("--source-url",help="Official retrieval link for IPEDS source file")
    parser.add_argument("--validate",action="store_true")
    parser.add_argument("--import",dest="do_import",action="store_true")
    parser.add_argument("--confirm-project",default="")
    a=parser.parse_args()
    if a.validate == a.do_import: raise ValueError("Choose exactly one of --validate or --import")
    if a.source=="bls":
        blob=Path(a.file).read_bytes() if a.file else fetch(BLS_URL)
        digest=sha256(blob)
        if digest!=BLS_SHA256: raise ValueError("Unreviewed BLS file checksum changed; ingestion stopped")
        rows=parse_bls(blob)
        release="BLS_EMP_2025_2035"
        src=BLS_URL
    else:
        if not a.file or not a.source_url or not a.source_url.startswith("https://"):
            raise ValueError("IPEDS requires --file of downloaded C2025_A ZIP and verified --source-url")
        blob=Path(a.file).read_bytes()
        if len(blob)>MAX_BYTES: raise ValueError("IPEDS file exceeds size cap")
        rows=parse_ipeds(blob,a.source_url)
        release="C2025_A_provisional"
        src=a.source_url
    print(json.dumps({
        "source_release":release,"source_sha256":sha256(blob),
        "rows":len(rows),"codes":len({r["taxonomy_code"] for r in rows}),
        "source_bytes":len(blob),"minimum":min(Decimal(r["metric_value"]) for r in rows),
        "maximum":max(Decimal(r["metric_value"]) for r in rows),
        "notice":"Raw observations stay private; no ranking of licenses or certifications inferred."
    },default=str,indent=2))
    if a.do_import:
        ingest(rows,src,sha256(blob),len(blob),a.confirm_project,release)

if __name__=="__main__":
    try:
        main()
    except (ValueError,RuntimeError,urllib.error.URLError,KeyError,IndexError) as e:
        print(f"Official demand ingestion failed closed: {e}",file=sys.stderr)
        sys.exit(2)
